from __future__ import annotations

import base64
import csv
import io
import json
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import yaml
from docx import Document as DocxDocument
from pypdf import PdfReader

from ..builder import SourceDocument

TEXT_EXTENSIONS = {
    ".md",
    ".mdx",
    ".txt",
    ".rst",
    ".csv",
    ".json",
    ".yaml",
    ".yml",
    ".toml",
    ".py",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".html",
    ".htm",
    ".xml",
    ".log",
    ".eml",
}

MAX_DOCUMENT_CHARS = 120_000


def parse_uploaded_document(
    *,
    name: str,
    text: str | None = None,
    content_base64: str | None = None,
    mime_type: str | None = None,
) -> SourceDocument:
    """Parse an uploaded browser document into text for brain ingestion."""
    suffix = Path(name).suffix.lower()
    raw = _decode_base64(content_base64) if content_base64 else None

    if text and not raw:
        parsed = text
    elif raw is not None:
        parsed = _parse_bytes(name=name, suffix=suffix, data=raw, mime_type=mime_type or "")
    else:
        raise ValueError(f"{name} has no readable text or binary content")

    parsed = parsed.strip()
    if not parsed:
        raise ValueError(f"{name} did not contain extractable text")
    return SourceDocument(name=name, text=parsed[:MAX_DOCUMENT_CHARS], source_path=name)


def _decode_base64(content_base64: str | None) -> bytes:
    if not content_base64:
        return b""
    _, _, payload = content_base64.partition(",")
    body = payload or content_base64
    return base64.b64decode(body)


def _parse_bytes(*, name: str, suffix: str, data: bytes, mime_type: str) -> str:
    if suffix == ".pdf" or mime_type == "application/pdf":
        return _parse_pdf(data)
    if suffix == ".docx" or mime_type.endswith("wordprocessingml.document"):
        return _parse_docx(data)
    if suffix == ".pptx" or mime_type.endswith("presentationml.presentation"):
        return _parse_office_zip_xml(data, label="slide")
    if suffix == ".xlsx" or mime_type.endswith("spreadsheetml.sheet"):
        return _parse_office_zip_xml(data, label="sheet")
    if suffix in {".json"}:
        return json.dumps(json.loads(_decode_text(data)), indent=2, ensure_ascii=True)
    if suffix in {".yaml", ".yml"}:
        loaded = yaml.safe_load(_decode_text(data))
        return yaml.safe_dump(loaded, sort_keys=False, allow_unicode=False)
    if suffix == ".csv":
        return _parse_csv(data)
    if suffix in TEXT_EXTENSIONS or mime_type.startswith("text/"):
        return _decode_text(data)
    return _decode_text(data)


def _parse_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    chunks: list[str] = []
    for index, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            chunks.append(f"\n\n--- PDF PAGE {index} ---\n{text.strip()}")
    return "\n".join(chunks)


def _parse_docx(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    chunks: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            chunks.append(para.text.strip())
    for table_index, table in enumerate(doc.tables, 1):
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows.append(" | ".join(cells))
        if rows:
            chunks.append(f"\n--- TABLE {table_index} ---\n" + "\n".join(rows))
    return "\n\n".join(chunks)


def _parse_office_zip_xml(data: bytes, *, label: str) -> str:
    chunks: list[str] = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        xml_names = [
            name
            for name in zf.namelist()
            if name.endswith(".xml") and ("/slides/" in name or "/worksheets/" in name or "/sharedStrings" in name)
        ]
        for index, xml_name in enumerate(xml_names, 1):
            try:
                root = ElementTree.fromstring(zf.read(xml_name))
            except ElementTree.ParseError:
                continue
            texts = [
                (node.text or "").strip()
                for node in root.iter()
                if node.tag.endswith("}t") and (node.text or "").strip()
            ]
            if texts:
                chunks.append(f"\n--- {label.upper()} {index}: {xml_name} ---\n" + "\n".join(texts))
    return "\n".join(chunks)


def _parse_csv(data: bytes) -> str:
    text = _decode_text(data)
    rows = list(csv.reader(io.StringIO(text)))
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")
